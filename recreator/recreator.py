import os
import json
import logging
from pptx import Presentation
from recreator.shape_recreators.text_shape_recreator import TextShapeRecreator
from recreator.shape_recreators.table_shape_recreator import TableShapeRecreator
from recreator.shape_recreators.picture_shape_recreator import PictureShapeRecreator
from recreator.shape_recreators.diagram_shape_recreator import DiagramShapeRecreator
from pptx.dml.color import RGBColor
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

class Recreator:
    def __init__(self, info_path, template_pptx_path=None):
        self.info_path = info_path
        self.template_pptx_path = template_pptx_path
        self.recreators = [
            TextShapeRecreator(),
            TableShapeRecreator(),
            DiagramShapeRecreator(),
            PictureShapeRecreator()
        ]

    def _parse_slide_info(self, slide_info):
        presentation = {}
        slides = slide_info
        if isinstance(slide_info, dict):
            presentation = slide_info.get("presentation") or {}
            slides = slide_info.get("slides") or []
        return presentation, slides

    def _shape_geometry(self, shape):
        try:
            return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))
        except Exception:
            return None

    def _shape_data_geometry(self, shape_data):
        for key in ("left", "top", "width", "height"):
            if key not in shape_data:
                return None
        try:
            return (
                int(shape_data["left"]),
                int(shape_data["top"]),
                int(shape_data["width"]),
                int(shape_data["height"]),
            )
        except Exception:
            return None

    def _best_match_shape(self, slide, shape_data):
        desired = self._shape_data_geometry(shape_data)
        if not desired:
            return None

        shape_type = shape_data.get("type")
        candidates = []
        for shape in slide.shapes:
            if shape_type == "text" and not getattr(shape, "has_text_frame", False):
                continue
            if shape_type == "table" and not getattr(shape, "has_table", False):
                continue
            if shape_type == "picture":
                try:
                    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                        continue
                except Exception:
                    continue

            geo = self._shape_geometry(shape)
            if not geo:
                continue
            distance = sum(abs(a - b) for a, b in zip(desired, geo))
            candidates.append((distance, shape))

        if not candidates:
            return None

        candidates.sort(key=lambda x: x[0])
        distance, best = candidates[0]
        # Extracted-from-template should match exactly; allow small tolerance for rounding.
        if distance <= 5000:
            return best
        return None

    def recreate_pptx(self, output_dir, output_path=None):
        with open(self.info_path, "r") as f:
            slide_info = json.load(f)

        presentation, slides = self._parse_slide_info(slide_info)

        template_path = self.template_pptx_path
        if template_path and os.path.isfile(template_path) and template_path.lower().endswith(".pptx"):
            prs = Presentation(template_path)
            template_slides = list(prs.slides)
            for slide_index, slide_data in enumerate(slides):
                if slide_index < len(template_slides):
                    slide = template_slides[slide_index]
                else:
                    slide = prs.slides.add_slide(prs.slide_layouts[6])

                shapes_data = sorted(slide_data.get("shapes") or [], key=lambda x: x.get("z_order", 0))
                for shape_data in shapes_data:
                    if shape_data.get("type") == "diagram":
                        # Template PPTX already contains SmartArt; avoid duplicating it with text placeholders.
                        continue
                    for recreator in self.recreators:
                        if not recreator.can_recreate(shape_data):
                            continue
                        existing = self._best_match_shape(slide, shape_data)
                        if existing is not None and hasattr(recreator, "apply_to_shape"):
                            try:
                                if recreator.apply_to_shape(existing, shape_data):
                                    break
                            except Exception:
                                pass
                        recreator.recreate(slide, shape_data)
                        break
        else:
            prs = Presentation()
            slide_width = presentation.get("slide_width")
            slide_height = presentation.get("slide_height")
            if slide_width is not None:
                try:
                    prs.slide_width = Emu(int(slide_width))
                except (TypeError, ValueError):
                    pass
            if slide_height is not None:
                try:
                    prs.slide_height = Emu(int(slide_height))
                except (TypeError, ValueError):
                    pass

            for slide_data in slides:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

                shapes_data = sorted(slide_data.get("shapes") or [], key=lambda x: x.get("z_order", 0))
                for shape_data in shapes_data:
                    for recreator in self.recreators:
                        if recreator.can_recreate(shape_data):
                            recreator.recreate(slide, shape_data)
                            break

        os.makedirs(output_dir, exist_ok=True)
        if output_path:
            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        else:
            info_base = os.path.basename(self.info_path)
            if info_base.lower().endswith("_info.json"):
                stem = info_base[: -len("_info.json")]
            else:
                stem = os.path.splitext(info_base)[0]
            output_path = os.path.join(output_dir, f"recreated_{stem}.pptx")
        prs.save(output_path)

        return output_path

from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR

class TextShapeRecreator(BaseShapeRecreator):
    def can_recreate(self, shape_data):
        return shape_data["type"] == "text"

    def _emu_or_default(self, value, default):
        if value is None:
            return default
        if isinstance(value, (int, float)):
            try:
                return Emu(int(value))
            except (TypeError, ValueError):
                return default
        return value

    def _normalize_rgb(self, color):
        if not color:
            return None

        if isinstance(color, str):
            candidate = color.strip().lstrip("#").upper()
            if len(candidate) == 6 and all(c in "0123456789ABCDEF" for c in candidate):
                return candidate
            return None

        if isinstance(color, (bytes, bytearray)) and len(color) == 3:
            return color.hex().upper()

        if isinstance(color, (list, tuple)) and len(color) == 3:
            try:
                r, g, b = (int(color[0]), int(color[1]), int(color[2]))
            except (TypeError, ValueError):
                return None
            if 0 <= r <= 255 and 0 <= g <= 255 and 0 <= b <= 255:
                return f"{r:02X}{g:02X}{b:02X}"
            return None

        candidate = str(color).strip().lstrip("#").upper()
        if len(candidate) == 6 and all(c in "0123456789ABCDEF" for c in candidate):
            return candidate
        return None

    def apply_to_shape(self, shape, shape_data):
        if not getattr(shape, "has_text_frame", False):
            return False
        self._populate_text_frame(shape.text_frame, shape_data)
        return True

    def _populate_text_frame(self, text_frame, shape_data):
        tf = shape_data.get("text_frame") or {}
        vertical_anchor = tf.get("vertical_anchor")
        if vertical_anchor is not None:
            try:
                text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR(int(vertical_anchor))
            except (TypeError, ValueError):
                pass

        if tf.get("word_wrap") is not None:
            text_frame.word_wrap = bool(tf.get("word_wrap"))

        for margin_key, attr in (
            ("margin_left", "margin_left"),
            ("margin_right", "margin_right"),
            ("margin_top", "margin_top"),
            ("margin_bottom", "margin_bottom"),
        ):
            margin = tf.get(margin_key)
            if margin is not None:
                try:
                    setattr(text_frame, attr, Emu(int(margin)))
                except (TypeError, ValueError):
                    pass

        paragraphs_data = shape_data.get("paragraphs") or []
        if not paragraphs_data:
            return

        # Clear and rebuild content.
        try:
            text_frame.clear()
        except Exception:
            # Fallback for older python-pptx APIs
            for p in list(text_frame.paragraphs)[1:]:
                try:
                    p._p.getparent().remove(p._p)
                except Exception:
                    pass
            try:
                text_frame.paragraphs[0].clear()
            except Exception:
                pass

        paragraph = text_frame.paragraphs[0]
        for i, paragraph_data in enumerate(paragraphs_data):
            if i > 0:
                paragraph = text_frame.add_paragraph()

            alignment = paragraph_data.get("alignment")
            if alignment is not None:
                try:
                    paragraph.alignment = PP_ALIGN(int(alignment))
                except (TypeError, ValueError):
                    pass

            level = paragraph_data.get("level")
            if level is not None:
                try:
                    paragraph.level = int(level)
                except (TypeError, ValueError):
                    pass

            line_spacing = paragraph_data.get("line_spacing")
            if line_spacing is not None:
                try:
                    paragraph.line_spacing = Pt(float(line_spacing))
                except (TypeError, ValueError):
                    pass

            space_before = paragraph_data.get("space_before")
            if space_before is not None:
                try:
                    paragraph.space_before = Pt(float(space_before))
                except (TypeError, ValueError):
                    pass

            space_after = paragraph_data.get("space_after")
            if space_after is not None:
                try:
                    paragraph.space_after = Pt(float(space_after))
                except (TypeError, ValueError):
                    pass

            paragraph_font = paragraph_data.get("font") or {}
            if paragraph_font.get("name"):
                paragraph.font.name = paragraph_font.get("name")
            if paragraph_font.get("size") is not None:
                try:
                    paragraph.font.size = Pt(float(paragraph_font.get("size")))
                except (TypeError, ValueError):
                    pass
            paragraph.font.bold = paragraph_font.get("bold")
            paragraph.font.italic = paragraph_font.get("italic")
            paragraph.font.underline = paragraph_font.get("underline")
            paragraph_rgb = self._normalize_rgb(paragraph_font.get("color"))
            if paragraph_rgb:
                paragraph.font.color.rgb = RGBColor.from_string(paragraph_rgb)

            for run_data in paragraph_data.get("runs") or []:
                run = paragraph.add_run()
                run.text = run_data.get("text") or ""

                font_name = run_data.get("font")
                if font_name:
                    run.font.name = font_name

                font_size = run_data.get("size")
                if font_size is not None:
                    try:
                        run.font.size = Pt(float(font_size))
                    except (TypeError, ValueError):
                        pass

                run.font.bold = run_data.get("bold")
                run.font.italic = run_data.get("italic")
                run.font.underline = run_data.get("underline")

                rgb = self._normalize_rgb(run_data.get("color"))
                if rgb:
                    run.font.color.rgb = RGBColor.from_string(rgb)

    def recreate(self, slide, shape_data):
        left = self._emu_or_default(shape_data.get("left"), Inches(1))
        top = self._emu_or_default(shape_data.get("top"), Inches(1))
        width = self._emu_or_default(shape_data.get("width"), Inches(5))
        height = self._emu_or_default(shape_data.get("height"), Inches(1))

        text_box = slide.shapes.add_textbox(left, top, width, height)
        self._populate_text_frame(text_box.text_frame, shape_data)

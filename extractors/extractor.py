import os
import json
import logging
from pptx import Presentation
from extractors.shape_extractors.text_shape_extractor import TextShapeExtractor
from extractors.shape_extractors.table_shape_extractor import TableShapeExtractor
from extractors.shape_extractors.picture_shape_extractor import PictureShapeExtractor
from extractors.shape_extractors.diagram_shape_extractor import DiagramShapeExtractor

class Extractor:
    def __init__(self, pptx_path):
        self.pptx_path = pptx_path
        self.extractors = [
            TextShapeExtractor(),
            TableShapeExtractor(),
            DiagramShapeExtractor(),
            PictureShapeExtractor()
        ]

    def _extract_geometry(self, shape):
        geometry = {}
        for key in ("left", "top", "width", "height"):
            value = getattr(shape, key, None)
            if value is None:
                continue
            try:
                geometry[key] = int(value)
            except (TypeError, ValueError):
                continue

        rotation = getattr(shape, "rotation", None)
        if rotation is not None:
            try:
                geometry["rotation"] = float(rotation)
            except (TypeError, ValueError):
                pass

        return geometry

    def extract_info(self, output_dir):
        prs = Presentation(self.pptx_path)
        slide_info = {
            "version": 2,
            "source_pptx": {
                "path": self.pptx_path,
                "basename": os.path.basename(self.pptx_path),
            },
            "presentation": {
                "slide_width": int(prs.slide_width),
                "slide_height": int(prs.slide_height),
            },
            "slides": [],
        }

        for slide_index, slide in enumerate(prs.slides):
            slide_data = {"slide_index": slide_index, "shapes": []}
            for z_order, shape in enumerate(slide.shapes):
                for extractor in self.extractors:
                    if extractor.can_extract(shape):
                        shape_data = extractor.extract(shape)
                        shape_data["z_order"] = z_order
                        shape_data.update(self._extract_geometry(shape))
                        slide_data["shapes"].append(shape_data)
                        break
            slide_info["slides"].append(slide_data)

        os.makedirs(output_dir, exist_ok=True)  # Create the output directory if it doesn't exist
        output_path = os.path.join(output_dir, f"{os.path.basename(self.pptx_path).split('.')[0]}_info.json")
        with open(output_path, "w") as f:
            json.dump(slide_info, f, indent=2)

        return output_path

from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator  
from pptx.util import Inches, Pt, Emu  
from pptx.dml.color import RGBColor  
from pptx.enum.text import PP_ALIGN  
from pptx.table import Table  
from pptx.oxml.ns import qn  
from pptx.oxml.xmlchemy import OxmlElement  
  
class ExtendedRGBColor(RGBColor):  
    def __init__(self, rgb):  
        super().__init__(rgb)  
        self.r = (rgb >> 16) & 0xFF  
        self.g = (rgb >> 8) & 0xFF  
        self.b = rgb & 0xFF  
  
class TableShapeRecreator(BaseShapeRecreator):  
    def can_recreate(self, shape_data):  
        return shape_data["type"] == "table"  

    def apply_to_shape(self, shape, shape_data):
        if not getattr(shape, "has_table", False):
            return False

        table = shape.table

        col_widths = shape_data.get("col_widths") or []
        for i, col_width in enumerate(col_widths):
            if col_width is None:
                continue
            try:
                table.columns[i].width = Emu(int(col_width))
            except (IndexError, TypeError, ValueError):
                pass

        row_heights = shape_data.get("row_heights") or []
        for i, row_height in enumerate(row_heights):
            if row_height is None:
                continue
            try:
                table.rows[i].height = Emu(int(row_height))
            except (IndexError, TypeError, ValueError):
                pass

        rows_data = shape_data.get("rows") or []
        for i, row_data in enumerate(rows_data):
            for j, cell_data in enumerate(row_data):
                try:
                    cell = table.cell(i, j)
                except IndexError:
                    continue
                cell.text = (cell_data or {}).get("text", "")

        return True
  
    def recreate(self, slide, shape_data):  
        rows_count = len(shape_data["rows"])  
        cols_count = len(shape_data["rows"][0])  
  
        left = Emu(int(shape_data["left"])) if isinstance(shape_data.get("left"), (int, float)) else shape_data.get("left", Inches(1))
        top = Emu(int(shape_data["top"])) if isinstance(shape_data.get("top"), (int, float)) else shape_data.get("top", Inches(1))
        width = Emu(int(shape_data["width"])) if isinstance(shape_data.get("width"), (int, float)) else shape_data.get("width", Inches(6))
        height = Emu(int(shape_data["height"])) if isinstance(shape_data.get("height"), (int, float)) else shape_data.get("height", Inches(0.5) * rows_count)
  
        table = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table  

        col_widths = shape_data.get("col_widths") or []
        for i, col_width in enumerate(col_widths):
            if col_width is None:
                continue
            try:
                table.columns[i].width = Emu(int(col_width))
            except (IndexError, TypeError, ValueError):
                pass

        row_heights = shape_data.get("row_heights") or []
        for i, row_height in enumerate(row_heights):
            if row_height is None:
                continue
            try:
                table.rows[i].height = Emu(int(row_height))
            except (IndexError, TypeError, ValueError):
                pass
  
        for i, row_data in enumerate(shape_data["rows"]):  
            for j, cell_data in enumerate(row_data):  
                cell = table.cell(i, j)  
                cell.text = cell_data["text"]  
                cell.text_frame.paragraphs[0].font.size = Pt(10)  
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER  
                cell.fill.solid()  
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)  
  
                # Set border properties for the cell  
                def rgb_to_hex(color):  
                    return '{:02x}{:02x}{:02x}'.format(color[0], color[1], color[2])  
                
                border_color = RGBColor(0, 0, 0)  # black color  
                border_color_hex = rgb_to_hex(border_color)  
                border_width = Pt(1)  

  
  
                def set_cell_border(cell, border_side):  
                    tc = cell._tc  
                    tcPr = tc.get_or_add_tcPr()  

                    side_tag = {
                        "top": "a:lnT",
                        "right": "a:lnR",
                        "bottom": "a:lnB",
                        "left": "a:lnL",
                    }.get(border_side)
                    if not side_tag:
                        return

                    ln = tcPr.find(qn(side_tag))
                    if ln is None:  
                        ln = OxmlElement(side_tag)
                        tcPr.append(ln)  
  
                    ln.set('w', str(border_width.emu))
                    ln.set('cap', 'flat')  
                    ln.set('cmpd', 'sng')  
                    ln.set('algn', 'ctr')  
  
                    solidFill = ln.find(qn('a:solidFill'))  
                    if solidFill is None:  
                        solidFill = OxmlElement('a:solidFill')  
                        ln.append(solidFill)  
  
                    srgbClr = solidFill.find(qn('a:srgbClr'))  
                    if srgbClr is None:  
                        srgbClr = OxmlElement('a:srgbClr')  
                        solidFill.append(srgbClr)  
  
                    srgbClr.set('val', border_color_hex)  
  
                set_cell_border(cell, "top")  
                set_cell_border(cell, "right")  
                set_cell_border(cell, "bottom")  
                set_cell_border(cell, "left")  
  
                if cell_data["row_span"] > 1:  
                    cell.merge(table.cell(i + cell_data["row_span"] - 1, j))  
                if cell_data["col_span"] > 1:  
                    cell.merge(table.cell(i, j + cell_data["col_span"] - 1))  

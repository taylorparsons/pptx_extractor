import os
import json
import logging
from pptx import Presentation
from extractors.shape_extractors.text_shape_extractor import TextShapeExtractor
from extractors.shape_extractors.table_shape_extractor import TableShapeExtractor
from extractors.shape_extractors.picture_shape_extractor import PictureShapeExtractor

class Extractor:
    def __init__(self, pptx_path):
        self.pptx_path = pptx_path
        self.extractors = [
            TextShapeExtractor(),
            TableShapeExtractor(),
            PictureShapeExtractor()
        ]

    def extract_info(self, output_dir):
        prs = Presentation(self.pptx_path)
        slide_info = []

        for slide in prs.slides:
            slide_data = {"shapes": []}
            for shape in slide.shapes:
                for extractor in self.extractors:
                    if extractor.can_extract(shape):
                        shape_data = extractor.extract(shape)
                        slide_data["shapes"].append(shape_data)
                        break
            slide_info.append(slide_data)

        os.makedirs(output_dir, exist_ok=True)  # Create the output directory if it doesn't exist
        output_path = os.path.join(output_dir, f"{os.path.basename(self.pptx_path).split('.')[0]}_info.json")
        with open(output_path, "w") as f:
            json.dump(slide_info, f, indent=2)

        return output_path



from abc import ABC, abstractmethod

class BaseShapeExtractor(ABC):
    @abstractmethod
    def can_extract(self, shape):
        pass

    @abstractmethod
    def extract(self, shape):
        pass

from extractors.shape_extractors.base_shape_extractor import BaseShapeExtractor

class TextShapeExtractor(BaseShapeExtractor):
    def can_extract(self, shape):
        return shape.has_text_frame

    def extract(self, shape):
        text_frame = shape.text_frame
        paragraphs = []

        for paragraph in text_frame.paragraphs:
            paragraph_data = {"runs": []}
            for run in paragraph.runs:
                run_data = {
                    "text": run.text,
                    "font": run.font.name,
                    "size": run.font.size.pt if run.font.size else None,
                    "bold": run.font.bold,
                    "italic": run.font.italic,
                    "underline": run.font.underline,
                    "color": self.get_color(run.font.color)
                }
                paragraph_data["runs"].append(run_data)
            paragraphs.append(paragraph_data)

        return {
            "type": "text",
            "paragraphs": paragraphs
        }

    def get_color(self, color):
        if hasattr(color, 'rgb'):
            return color.rgb
        else:
            return None

from extractors.shape_extractors.base_shape_extractor import BaseShapeExtractor

class TableShapeExtractor(BaseShapeExtractor):
    def can_extract(self, shape):
        return shape.has_table

    def extract(self, shape):
        table = shape.table
        rows = []

        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_data = {
                    "text": cell.text,
                    "row_span": cell.span_height,
                    "col_span": cell.span_width
                }
                cells.append(cell_data)
            rows.append(cells)

        return {
            "type": "table",
            "rows": rows
        }

import base64
from extractors.shape_extractors.base_shape_extractor import BaseShapeExtractor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from io import BytesIO

class PictureShapeExtractor(BaseShapeExtractor):
    def can_extract(self, shape):
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE

    def extract(self, shape):
        pic = shape
        image_data = BytesIO(pic.image.blob)
        base64_image = base64.b64encode(image_data.getvalue()).decode('utf-8')
        return {
            "type": "picture",
            "image_data": base64_image
        }



import os
import json
import logging
from pptx import Presentation
from recreator.shape_recreators.text_shape_recreator import TextShapeRecreator
from recreator.shape_recreators.table_shape_recreator import TableShapeRecreator
from recreator.shape_recreators.picture_shape_recreator import PictureShapeRecreator
from pptx.dml.color import RGBColor

class Recreator:
    def __init__(self, info_path):
        self.info_path = info_path
        self.recreators = [
            TextShapeRecreator(),
            TableShapeRecreator(),
            PictureShapeRecreator()
        ]

    def recreate_pptx(self, output_dir):
        with open(self.info_path, "r") as f:
            slide_info = json.load(f)

        prs = Presentation()

        for slide_data in slide_info:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

            shapes_data = sorted(slide_data["shapes"], key=lambda x: x.get("z_order", 0))
            for shape_data in shapes_data:
                for recreator in self.recreators:
                    if recreator.can_recreate(shape_data):
                        recreator.recreate(slide, shape_data)
                        break

        output_path = os.path.join(output_dir, f"recreated_{os.path.basename(self.info_path).split('_')[0]}.pptx")
        prs.save(output_path)

        return output_path



from abc import ABC, abstractmethod

class BaseShapeRecreator(ABC):
    @abstractmethod
    def can_recreate(self, shape_data):
        pass

    @abstractmethod
    def recreate(self, slide, shape_data):
        pass

from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR

class TextShapeRecreator(BaseShapeRecreator):
    def can_recreate(self, shape_data):
        return shape_data["type"] == "text"

    def recreate(self, slide, shape_data):
        left = shape_data.get("left", Inches(1))
        top = shape_data.get("top", Inches(1))
        width = shape_data.get("width", Inches(5))
        height = shape_data.get("height", Inches(1))

        text_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = text_box.text_frame
        text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

        for paragraph_data in shape_data["paragraphs"]:
            paragraph = text_frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.line_spacing = Pt(12)
            for run_data in paragraph_data["runs"]:
                run = paragraph.add_run()
                run.text = run_data["text"]
                run.font.name = run_data["font"]
                run.font.size = Pt(run_data["size"])
                run.font.bold = run_data["bold"]
                run.font.italic = run_data["italic"]
                run.font.underline = run_data["underline"]
                if run_data["color"]:
                    run.font.color.rgb = RGBColor.from_string(run_data["color"])

from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator  
from pptx.util import Inches, Pt  
from pptx.dml.color import RGBColor  
from pptx.enum.text import PP_ALIGN  
from pptx.table import Table  
from pptx.oxml.ns import qn  
from pptx.oxml.xmlchemy import OxmlElement  
  
class ExtendedRGBColor(RGBColor):  
    def __init__(self, rgb):  
        super().__init__(rgb)  
        self.r = (rgb >> 16) & 0xFF  
        self.g = (rgb >> 8) & 0xFF  
        self.b = rgb & 0xFF  
  
class TableShapeRecreator(BaseShapeRecreator):  
    def can_recreate(self, shape_data):  
        return shape_data["type"] == "table"  
  
    def recreate(self, slide, shape_data):  
        rows_count = len(shape_data["rows"])  
        cols_count = len(shape_data["rows"][0])  
  
        left = shape_data.get("left", Inches(1))  
        top = shape_data.get("top", Inches(1))  
        width = shape_data.get("width", Inches(6))  
        height = shape_data.get("height", Inches(0.5) * rows_count)  
  
        table = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table  
  
        for i, row_data in enumerate(shape_data["rows"]):  
            for j, cell_data in enumerate(row_data):  
                cell = table.cell(i, j)  
                cell.text = cell_data["text"]  
                cell.text_frame.paragraphs[0].font.size = Pt(10)  
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER  
                cell.fill.solid()  
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)  
  
                # Set border properties for the cell  
                def rgb_to_hex(color):  
                    return '{:02x}{:02x}{:02x}'.format(color[0], color[1], color[2])  
                
                border_color = RGBColor(0, 0, 0)  # black color  
                border_color_hex = rgb_to_hex(border_color)  
                border_width = Pt(1)  

  
  
                def set_cell_border(cell, border_side):  
                    tc = cell._tc  
                    tcPr = tc.get_or_add_tcPr()  
  
                    lnL = tcPr.find(qn('a:lnL'))  
                    if lnL is None:  
                        lnL = OxmlElement('a:lnL')  
                        tcPr.append(lnL)  
  
                    lnL.set('w', str(border_width.centipoints))  
                    lnL.set('cap', 'flat')  
                    lnL.set('cmpd', 'sng')  
                    lnL.set('algn', 'ctr')  
  
                    solidFill = lnL.find(qn('a:solidFill'))  
                    if solidFill is None:  
                        solidFill = OxmlElement('a:solidFill')  
                        lnL.append(solidFill)  
  
                    srgbClr = solidFill.find(qn('a:srgbClr'))  
                    if srgbClr is None:  
                        srgbClr = OxmlElement('a:srgbClr')  
                        solidFill.append(srgbClr)  
  
                    srgbClr.set('val', border_color_hex)  
  
                set_cell_border(cell, "top")  
                set_cell_border(cell, "right")  
                set_cell_border(cell, "bottom")  
                set_cell_border(cell, "left")  
  
                if cell_data["row_span"] > 1:  
                    cell.merge(table.cell(i + cell_data["row_span"] - 1, j))  
                if cell_data["col_span"] > 1:  
                    cell.merge(table.cell(i, j + cell_data["col_span"] - 1))  


import base64
from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator
from pptx.util import Inches
from io import BytesIO

class PictureShapeRecreator(BaseShapeRecreator):
    def can_recreate(self, shape_data):
        return shape_data["type"] == "picture"

    def recreate(self, slide, shape_data):
        left = shape_data.get("left", Inches(1))
        top = shape_data.get("top", Inches(1))
        width = shape_data.get("width", Inches(1))
        height = shape_data.get("height", Inches(1))
        base64_image = shape_data["image_data"]
        image_data = base64.b64decode(base64_image)
        pic = BytesIO(image_data)
        slide.shapes.add_picture(pic, left, top, width, height)



import argparse

def parse_arguments():
    parser = argparse.ArgumentParser(description="PPTX Utility Tool")
    parser.add_argument("pptx_path", help="Path to the PPTX file to be processed.")
    parser.add_argument("output_dir", help="Directory to save the output files.")
    parser.add_argument("--extract", action="store_true", help="Extract presentation information.")
    parser.add_argument("--recreate", action="store_true", help="Recreate PPTX from extracted information.")
    return parser.parse_args()

def rgb_to_hex(rgb):
    if isinstance(rgb, tuple) and len(rgb) == 3:
        return '#{:02x}{:02x}{:02x}'.format(*rgb)
    return None

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    if len(hex_color) == 6:
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    return None

import os
import logging
from utils.argument_parser import parse_arguments
from extractors.extractor import Extractor
from recreator.recreator import Recreator

def main():
    args = parse_arguments()

    if args.extract:
        extractor = Extractor(args.pptx_path)
        info_path = extractor.extract_info(args.output_dir)
        logging.info(f"Presentation information saved to: {info_path}")

    if args.recreate:
        info_path = os.path.join(args.output_dir, f"{os.path.basename(args.pptx_path).split('.')[0]}_info.json")
        if not os.path.exists(info_path):
            raise FileNotFoundError(f"Info file not found: {info_path}")
        recreator = Recreator(info_path)
        output_pptx_path = recreator.recreate_pptx(args.output_dir)
        logging.info(f"New PPTX created at: {output_pptx_path}")

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
    try:
        main()
    except Exception as e:
        logging.error(f"Error in main execution: {str(e)}")
        raise

# PPTX Utility Tool
README_TEXT = """
The PPTX Utility Tool is a Python project that allows you to extract presentation information from a PowerPoint file (PPTX) and recreate the PPTX file using the extracted information. This tool can be useful for analyzing and manipulating PowerPoint presentations programmatically.

## Prerequisites

Before using this tool, make sure you have the following:

- Python 3.x installed on your system
- `python-pptx` library installed (`pip install python-pptx`)

## Getting Started

To get started with the PPTX Utility Tool, follow these steps:

1. Clone the repository or download the project files to your local machine.

2. Open a terminal or command prompt and navigate to the project directory.

3. Create a virtual environment to isolate the project dependencies:
   ```sh
   python3 -m venv env
   ```

4. Activate the virtual environment:
   ```sh
   source env/bin/activate
   ```

5. Install the required dependencies:
   ```sh
   pip install -r requirements.txt
   ```

## Usage

The PPTX Utility Tool provides two main functionalities: extracting presentation information and recreating the PPTX file.

### Extracting Presentation Information

To extract presentation information from a PPTX file, use the following command:

```sh
python main.py <pptx_path> <output_dir> --extract
```

- `<pptx_path>`: Path to the PPTX file you want to extract information from.
- `<output_dir>`: Directory where the extracted information will be saved as a JSON file.

Example:
```sh
python main.py presentation.pptx output --extract
```

This command will extract the presentation information from `presentation.pptx` and save it as a JSON file in the `output` directory.

### Recreating the PPTX File

To recreate the PPTX file using the extracted information, use the following command:

```sh
python main.py <pptx_path> <output_dir> --recreate
```

- `<pptx_path>`: Path to the PPTX file you want to recreate (used to determine the JSON file name).
- `<output_dir>`: Directory where the recreated PPTX file will be saved.

Example:
```sh
python main.py presentation.pptx output --recreate
```

This command will look for the extracted information JSON file (`presentation_info.json`) in the `output` directory and use it to recreate the PPTX file. The recreated file will be saved as `recreated_presentation.pptx` in the `output` directory.

## Troubleshooting

If you encounter any issues or errors while using the PPTX Utility Tool, please check the following:

- Make sure you have installed the required dependencies (`python-pptx`).
- Ensure that you are providing the correct paths for the PPTX file and output directory.
- Check the error messages logged in the console for any specific details about the issue.

## License

This project is licensed under the [MIT License](LICENSE).
"""

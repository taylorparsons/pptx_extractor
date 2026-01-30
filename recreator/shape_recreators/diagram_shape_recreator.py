from recreator.shape_recreators.base_shape_recreator import BaseShapeRecreator
from pptx.util import Inches, Emu, Pt


class DiagramShapeRecreator(BaseShapeRecreator):
    def can_recreate(self, shape_data):
        return shape_data.get("type") == "diagram"

    def _emu_or_default(self, value, default):
        if value is None:
            return default
        if isinstance(value, (int, float)):
            try:
                return Emu(int(value))
            except (TypeError, ValueError):
                return default
        return value

    def recreate(self, slide, shape_data):
        left = self._emu_or_default(shape_data.get("left"), Inches(1))
        top = self._emu_or_default(shape_data.get("top"), Inches(1))
        width = self._emu_or_default(shape_data.get("width"), Inches(5))
        height = self._emu_or_default(shape_data.get("height"), Inches(1))

        text_box = slide.shapes.add_textbox(left, top, width, height)

        rotation = shape_data.get("rotation")
        if rotation is not None:
            try:
                text_box.rotation = float(rotation)
            except (TypeError, ValueError):
                pass

        text_frame = text_box.text_frame
        text_frame.word_wrap = True

        texts = shape_data.get("texts") or []
        if not texts:
            return

        try:
            text_frame.clear()
        except Exception:
            pass

        paragraph = text_frame.paragraphs[0]
        paragraph.text = str(texts[0])
        try:
            paragraph.font.size = Pt(12)
        except Exception:
            pass

        for text in texts[1:]:
            p = text_frame.add_paragraph()
            p.text = str(text)
            try:
                p.font.size = Pt(12)
            except Exception:
                pass
